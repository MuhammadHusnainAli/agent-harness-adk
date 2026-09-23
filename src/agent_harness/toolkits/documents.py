"""Document parsing and OCR: get text out of whatever the user dropped in.

Text, Markdown, CSV, TSV, JSON, JSONL, HTML and XML parse with no dependencies
at all. PDF, DOCX and OCR need an optional library each; when one is missing the
tool says exactly which `pip install` unlocks that format instead of failing
with an import error halfway through a run.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

from ..errors import ToolError
from ..tools import Tool, tool

__all__ = ["parse_document", "parse_text", "supported_formats", "make_document_tool"]

_TAGS = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.DOTALL | re.IGNORECASE)
_MARKUP = re.compile(r"<[^>]+>")
_BLANKS = re.compile(r"\n{3,}")

NATIVE = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json",
          ".jsonl", ".ndjson", ".html", ".htm", ".xml", ".yaml", ".yml", ".py",
          ".js", ".ts", ".sql", ".sh", ".toml", ".ini", ".cfg"}
OPTIONAL = {
    ".pdf": ("pypdf", "pypdf"),
    ".docx": ("docx", "python-docx"),
    ".png": ("pytesseract", "pytesseract pillow"),
    ".jpg": ("pytesseract", "pytesseract pillow"),
    ".jpeg": ("pytesseract", "pytesseract pillow"),
    ".tif": ("pytesseract", "pytesseract pillow"),
    ".tiff": ("pytesseract", "pytesseract pillow"),
}


def supported_formats() -> dict[str, list[str]]:
    """Which formats work right now, and which need an install first."""
    available, missing = [], []
    for suffix, (module, install) in OPTIONAL.items():
        try:
            __import__(module)
            available.append(suffix)
        except ImportError:
            missing.append(f"{suffix} (pip install {install})")
    return {"native": sorted(NATIVE), "available": sorted(available),
            "needs_install": sorted(missing)}


def _html_to_text(body: str) -> str:
    stripped = _MARKUP.sub(" ", _TAGS.sub(" ", body))
    for entity, char in (("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"),
                         ("&gt;", ">"), ("&quot;", '"'), ("&#39;", "'")):
        stripped = stripped.replace(entity, char)
    return _BLANKS.sub("\n\n", re.sub(r"[ \t]{2,}", " ", stripped)).strip()


def _csv_to_text(body: str, delimiter: str = ",", *, max_rows: int = 500) -> str:
    rows = list(csv.reader(io.StringIO(body), delimiter=delimiter))
    if not rows:
        return ""
    header, *data = rows
    lines = [" | ".join(header)]
    for row in data[:max_rows]:
        lines.append(" | ".join(row))
    if len(data) > max_rows:
        lines.append(f"... [{len(data) - max_rows} more rows]")
    lines.append(f"\n({len(data)} rows, {len(header)} columns)")
    return "\n".join(lines)


def parse_text(content: str, suffix: str = ".txt") -> str:
    """Parse content already in memory. `suffix` says how to read it."""
    suffix = suffix.lower()
    if suffix in {".html", ".htm", ".xml"}:
        return _html_to_text(content)
    if suffix == ".csv":
        return _csv_to_text(content)
    if suffix == ".tsv":
        return _csv_to_text(content, "\t")
    if suffix == ".json":
        try:
            return json.dumps(json.loads(content), indent=2, ensure_ascii=False)
        except json.JSONDecodeError as exc:
            raise ToolError(f"invalid JSON: {exc}", tool="parse_document") from exc
    if suffix in {".jsonl", ".ndjson"}:
        out = []
        for index, line in enumerate(content.splitlines()):
            if line.strip():
                try:
                    out.append(json.dumps(json.loads(line), ensure_ascii=False))
                except json.JSONDecodeError:
                    out.append(f"[line {index + 1}: not valid JSON]")
        return "\n".join(out)
    return content


def _read_pdf(path: Path, *, max_pages: int) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ToolError(
            "reading PDFs needs pypdf — `pip install pypdf`", tool="parse_document"
        ) from exc
    reader = PdfReader(str(path))
    pages = reader.pages[:max_pages]
    text = "\n\n".join(f"--- page {i + 1} ---\n{(p.extract_text() or '').strip()}"
                       for i, p in enumerate(pages))
    meta = {"pages": len(reader.pages), "pages_read": len(pages)}
    if not text.strip():
        text = ("[no extractable text — this looks like a scanned PDF. Convert the "
                "pages to images and use OCR.]")
    return text, meta


def _read_docx(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import docx
    except ImportError as exc:
        raise ToolError(
            "reading .docx needs python-docx — `pip install python-docx`",
            tool="parse_document",
        ) from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), {"paragraphs": len(document.paragraphs),
                              "tables": len(document.tables)}


def _read_image(path: Path, *, language: str) -> tuple[str, dict[str, Any]]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ToolError(
            "OCR needs pytesseract and pillow — `pip install pytesseract pillow` "
            "(and the tesseract binary on PATH)", tool="parse_document",
        ) from exc
    with Image.open(path) as image:
        text = pytesseract.image_to_string(image, lang=language)
        return text.strip(), {"size": f"{image.width}x{image.height}", "ocr": True}


def parse_document(path: str | Path, *, max_chars: int = 200_000,
                   max_pages: int = 50, language: str = "eng") -> dict[str, Any]:
    """Read a document and return its text plus what we learned about it."""
    file = Path(path)
    if not file.is_file():
        raise ToolError(f"no such file: {path}", tool="parse_document")

    suffix = file.suffix.lower()
    meta: dict[str, Any] = {"path": str(file), "format": suffix or "unknown",
                            "bytes": file.stat().st_size}

    if suffix == ".pdf":
        text, extra = _read_pdf(file, max_pages=max_pages)
    elif suffix == ".docx":
        text, extra = _read_docx(file)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        text, extra = _read_image(file, language=language)
    elif suffix in NATIVE or not suffix:
        raw = file.read_text(encoding="utf-8", errors="replace")
        text, extra = parse_text(raw, suffix or ".txt"), {}
    else:
        formats = supported_formats()
        raise ToolError(
            f"cannot read {suffix or 'this file'}. Native: "
            f"{', '.join(sorted(NATIVE)[:8])}... "
            + (f"Needs an install: {'; '.join(formats['needs_install'])}"
               if formats["needs_install"] else ""),
            tool="parse_document",
        )

    meta.update(extra)
    meta["characters"] = len(text)
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n... [truncated at {max_chars} characters]"
        meta["truncated"] = True
    return {"text": text, **meta}


def make_document_tool(workspace: Any = None, *, name: str = "parse_document") -> Tool:
    """A document-parsing tool. With a workspace, paths resolve inside its jail."""

    @tool(name=name, tags=["builtin", "documents"], cacheable=True)
    def parse_document_tool(path: str, max_pages: int = 50) -> dict[str, Any]:
        """Read a document and return its text.

        Handles text, Markdown, CSV, TSV, JSON, JSONL, HTML and XML directly;
        PDF, DOCX and scanned images if the optional libraries are installed.

        Args:
            path: the file to read.
            max_pages: for PDFs, how many pages to read.
        """
        target = workspace.resolve(path) if workspace is not None else Path(path)
        return parse_document(target, max_pages=max_pages)

    return parse_document_tool
